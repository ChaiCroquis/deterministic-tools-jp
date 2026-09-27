"""excel_report のテスト。fixture は全て合成データ(fixtures/make_fixtures.py)、氏名も金額も項目も架空。

後半の test_measure_* は、合成した 1,000 行の差異一覧で
  (a) 差異を色だけで示した出力(理由コードの列も凡例も判断待ちシートも持たない)
  (b) 塗りを条件付き書式で指定した出力(開いた時に Excel が評価する形)
  (c) この部品(理由コードを値として置き、塗りを静的に焼き、読み戻して突き合わせる)
を比べる(記事の数値はここから取る。固定 seed で決定論的に再現し、変わったら記事の数字も変える)。
"""
from __future__ import annotations

import random
import re
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
FX = ROOT / "fixtures"
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(FX))
import excel_report as X  # noqa: E402
import make_fixtures as MF  # noqa: E402

REASONS = X.ReasonTable.from_csv(FX / "riyu_hyou.csv")
TABLE = X.Table.from_csv(FX / "sai_ichiran.csv")
JUDGMENTS = X._judgments_from_csv(FX / "hantei.csv")


def small() -> tuple[X.Table, list[X.Judgment]]:
    """そのまま書ける小さな表と判定 1 件。"""
    t = X.Table(("社員番号", "氏名", "差"), (("000001", "山田太郎", "0"), ("000002", "鈴木花子", "1525")))
    return t, [X.Judgment(2, "氏名", "相手が決まらない")]


# ---- 理由コードの表(呼ぶ側が渡す語彙) --------------------------------------------------------

def test_reason_table_from_csv():
    assert REASONS.codes == ("値が無い", "表の外にある", "相手が決まらない", "形が両立しない", "往復で値が変わる")
    assert REASONS.fill_of("値が無い") == "FFF1E8"
    assert REASONS.meaning_of("往復で値が変わる").startswith("書いて読み戻すと")


def test_reason_table_rejects_unknown_code():
    with pytest.raises(X.SpecError) as e:
        REASONS.fill_of("なにか別の理由")
    assert "表に無いコード" in str(e.value)


@pytest.mark.parametrize("kwargs", [
    {"code": "", "meaning": "意味", "fill": "FFF1E8"},          # コードが空
    {"code": "x", "meaning": "", "fill": "FFF1E8"},             # 凡例に書く 1 行説明が無い
    {"code": "x", "meaning": "意味", "fill": "fff1e8"},         # 小文字
    {"code": "x", "meaning": "意味", "fill": "FFF1E"},          # 5 桁
    {"code": "x", "meaning": "意味", "fill": "#FFF1E8"},        # # 付き
    {"code": "x", "meaning": "意味", "fill": "赤"},             # 色名
])
def test_reason_rejects_broken_row(kwargs):
    with pytest.raises(X.SpecError):
        X.Reason(**kwargs)


def test_reason_table_rejects_duplicate_code():
    with pytest.raises(X.SpecError):
        X.ReasonTable((X.Reason("x", "意味", "FFF1E8"), X.Reason("x", "別の意味", "E7ECFB")))


def test_reason_table_rejects_empty():
    with pytest.raises(X.SpecError):
        X.ReasonTable(())


def test_reason_table_needs_the_three_fields():
    with pytest.raises(X.SpecError) as e:
        X.ReasonTable.from_rows([{"理由コード": "x", "意味": "意味"}])
    assert "塗り" in str(e.value)


def test_two_codes_may_share_a_fill_but_the_legend_still_separates_them():
    """同じ淡色を 2 つのコードに割り当てても止めない(色は意味の運び手ではないので、表が正本)。"""
    r = X.ReasonTable((X.Reason("a", "意味 a", "FFF1E8"), X.Reason("b", "意味 b", "FFF1E8")))
    assert r.fill_of("a") == r.fill_of("b") == "FFF1E8"
    assert [x.meaning for x in r.reasons] == ["意味 a", "意味 b"]


# ---- 値の表 ----------------------------------------------------------------------------------

def test_table_from_csv():
    assert TABLE.columns == ("社員番号", "氏名", "項目", "前回", "今回", "差", "備考")
    assert len(TABLE.rows) == 12
    assert TABLE.value_at(1, "氏名") == "山田太郎"
    assert TABLE.value_at(3, "氏名") == ""


def test_table_rejects_a_column_named_like_the_reason_column():
    with pytest.raises(X.SpecError) as e:
        X.Table(("氏名", X.REASON_COLUMN), (("山田太郎", "値が無い"),))
    assert "衝突" in str(e.value)


def test_table_rejects_duplicate_heading():
    with pytest.raises(X.SpecError):
        X.Table(("氏名", "氏名"), (("a", "b"),))


def test_table_rejects_empty_columns():
    with pytest.raises(X.SpecError):
        X.Table((), ())


def test_table_rejects_ragged_row():
    with pytest.raises(X.SpecError) as e:
        X.Table(("a", "b"), (("1", "2"), ("1",)))
    assert "2 行目" in str(e.value)


def test_values_are_kept_as_strings():
    """先頭ゼロのコードも全角数字も、型を推測せずに文字列のまま置く。"""
    t = X.Table(("社員番号", "今回"), (("000042", "１９７００"),))
    assert t.value_at(1, "社員番号") == "000042" and t.value_at(1, "今回") == "１９７００"


# ---- 受け取ったものの検査(色を決める前に止まる) ----------------------------------------------

def test_check_judgments_passes_the_fixture():
    assert X.check_judgments(TABLE, REASONS, JUDGMENTS) is None


def test_unknown_reason_code_stops():
    with pytest.raises(X.SpecError) as e:
        X.check_judgments(TABLE, REASONS, [X.Judgment(1, "氏名", "なにか別の理由")])
    assert "表に無いコード" in str(e.value)


def test_row_outside_the_table_stops():
    with pytest.raises(X.SpecError) as e:
        X.check_judgments(TABLE, REASONS, [X.Judgment(13, "氏名", "値が無い")])
    assert "無い行" in str(e.value)


def test_row_zero_stops():
    with pytest.raises(X.SpecError):
        X.check_judgments(TABLE, REASONS, [X.Judgment(0, "氏名", "値が無い")])


def test_column_outside_the_table_stops():
    with pytest.raises(X.SpecError) as e:
        X.check_judgments(TABLE, REASONS, [X.Judgment(1, "存在しない列", "値が無い")])
    assert "無い列" in str(e.value)


def test_two_reasons_on_one_cell_stop():
    with pytest.raises(X.SpecError) as e:
        X.check_judgments(TABLE, REASONS, [X.Judgment(1, "氏名", "値が無い"),
                                           X.Judgment(1, "氏名", "相手が決まらない")])
    assert "同じセルに理由が 2 つ" in str(e.value)


def test_two_reasons_on_one_row_are_allowed():
    X.check_judgments(TABLE, REASONS, [X.Judgment(4, "前回", "値が無い"), X.Judgment(4, "差", "値が無い")])


# ---- 理由コードの列(値として置く側) ----------------------------------------------------------

def test_reason_column_is_placed_next_to_the_values():
    t, j = small()
    out = X.reason_column_values(t, j)
    assert out == ["", "相手が決まらない"]


def test_reason_column_joins_in_column_order():
    j = [X.Judgment(1, "差", "形が両立しない"), X.Judgment(1, "氏名", "値が無い")]
    t = X.Table(("氏名", "差"), (("", "1,525"),))
    assert X.reason_column_values(t, j) == ["値が無い 形が両立しない"]


def test_reason_column_writes_the_same_code_once():
    t = X.Table(("前回", "今回"), (("", ""),))
    j = [X.Judgment(1, "前回", "値が無い"), X.Judgment(1, "今回", "値が無い")]
    assert X.reason_column_values(t, j) == ["値が無い"]


def test_cell_coordinate_skips_the_header_row():
    assert X.cell_of(TABLE, 1, "社員番号") == "A2"
    assert X.cell_of(TABLE, 12, "備考") == "G13"


def test_counts_by_reason_keeps_zero_codes():
    got = X.counts_by_reason(REASONS, [X.Judgment(1, "氏名", "値が無い")])
    assert got == {"値が無い": 1, "表の外にある": 0, "相手が決まらない": 0,
                   "形が両立しない": 0, "往復で値が変わる": 0}


# ---- 書いて読み戻す --------------------------------------------------------------------------

def test_write_makes_three_sheets_in_a_fixed_order(tmp_path):
    out = tmp_path / "out.xlsx"
    assert X.write(out, TABLE, REASONS, JUDGMENTS).ok
    assert load_workbook(out).sheetnames == [X.SHEET_DETAIL, X.SHEET_LEGEND, X.SHEET_PENDING]


def test_write_returns_counts(tmp_path):
    r = X.write(tmp_path / "out.xlsx", TABLE, REASONS, JUDGMENTS)
    assert r.ok and r.rows == 12 and r.pending == 14
    assert r.by_reason == {"値が無い": 3, "表の外にある": 3, "相手が決まらない": 3,
                           "形が両立しない": 3, "往復で値が変わる": 2}


def test_detail_sheet_puts_the_values_as_they_were(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    got = X.read_values(out)
    assert got[0] == list(TABLE.columns) + [X.REASON_COLUMN]
    assert got[1][:len(TABLE.columns)] == list(TABLE.rows[0])
    assert got[1][0] == "000001"                      # 先頭ゼロが残る


def test_detail_sheet_has_the_reason_column_filled_in(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    got = X.read_values(out)
    by_row = {r[0]: r[-1] for r in got[1:]}
    assert by_row["000003"] == "値が無い"
    assert by_row["000005"] == "往復で値が変わる"
    assert by_row["000008"] == "相手が決まらない 表の外にある"       # 列の順で並ぶ
    assert by_row["000001"] == ""


def test_legend_sheet_carries_the_meaning_as_text(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    got = X.read_legend(out)
    assert [c for c, _m, _f in got] == list(REASONS.codes)
    assert all(m.strip() for _c, m, _f in got)
    assert [f for _c, _m, f in got] == ["FFF1E8", "E7ECFB", "EEF0F5", "FCEFC7", "EDE4F5"]


def test_legend_sheet_also_paints_a_swatch(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    assert X.read_fills(out, X.SHEET_LEGEND) == {"D2": "FFF1E8", "D3": "E7ECFB", "D4": "EEF0F5",
                                                 "D5": "FCEFC7", "D6": "EDE4F5"}


def test_pending_sheet_is_one_line_per_cell(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    rows = X.read_values(out, X.SHEET_PENDING)
    assert rows[0] == list(X.PENDING_FIELDS)
    assert len(rows) - 1 == 14
    assert rows[1] == [X.SHEET_DETAIL, "B4", "3", "氏名", "", "値が無い"]


def test_pending_sheet_is_sorted_by_row_then_column(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    rows = X.read_values(out, X.SHEET_PENDING)[1:]
    got = [(int(r[2]), r[3]) for r in rows]
    assert got == sorted(got, key=lambda p: (p[0], TABLE.index_of(p[1])))


def test_the_fill_follows_the_reason_table(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    fills = X.read_fills(out)
    assert len(fills) == 14
    for cell, code in X.expected_reason_map(TABLE, JUDGMENTS).items():
        assert fills[cell] == REASONS.fill_of(code)


def test_cells_without_a_reason_are_not_painted(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    painted = set(X.read_fills(out))
    assert "A2" not in painted and "G2" not in painted
    assert painted == set(X.expected_reason_map(TABLE, JUDGMENTS))


def test_the_reason_column_itself_is_not_painted(tmp_path):
    """理由コードの列は文字として読ませる欄なので塗らない(色は従属表示のまま)。"""
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    assert not [c for c in X.read_fills(out) if c.startswith("H")]     # H = 理由コードの列


def test_recover_reasons_reads_only_text(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    assert X.recover_reasons(out) == X.expected_reason_map(TABLE, JUDGMENTS)


def test_verify_passes_on_a_file_this_tool_wrote(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    assert X.verify(out, TABLE, REASONS, JUDGMENTS) == []


def test_write_is_deterministic_for_the_cells_and_fills(tmp_path):
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    X.write(a, TABLE, REASONS, JUDGMENTS)
    X.write(b, TABLE, REASONS, JUDGMENTS)
    assert X.read_values(a) == X.read_values(b)
    assert X.read_fills(a) == X.read_fills(b)
    assert X.recover_reasons(a) == X.recover_reasons(b)


# ---- 突合に差が出たら出力を消す ---------------------------------------------------------------

def test_a_repainted_cell_is_caught_and_the_reason_text_is_untouched(tmp_path):
    """塗りだけを別の理由の色に書き換えると、突合が塗りの差として 1 件返す。理由コードは動かない。"""
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    wb = load_workbook(out)
    wb[X.SHEET_DETAIL]["B4"].fill = X._solid(REASONS.fill_of("相手が決まらない"))
    wb.save(out)
    diffs = X.verify(out, TABLE, REASONS, JUDGMENTS)
    assert len(diffs) == 1 and diffs[0].what == "塗り" and diffs[0].cell == "B4"
    assert (diffs[0].original, diffs[0].readback) == ("FFF1E8", "EEF0F5")
    assert X.recover_reasons(out)["B4"] == "値が無い"


def test_an_extra_fill_on_a_clean_cell_is_caught(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    wb = load_workbook(out)
    wb[X.SHEET_DETAIL]["A2"].fill = X._solid("FFF1E8")
    wb.save(out)
    diffs = X.verify(out, TABLE, REASONS, JUDGMENTS)
    assert [(d.cell, d.what, d.original) for d in diffs] == [("A2", "塗り", "なし")]


def test_a_removed_fill_is_caught(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    wb = load_workbook(out)
    wb[X.SHEET_DETAIL]["B4"].fill = X.PatternFill()
    wb.save(out)
    diffs = X.verify(out, TABLE, REASONS, JUDGMENTS)
    assert [(d.cell, d.what, d.readback) for d in diffs] == [("B4", "塗り", "なし")]


def test_a_changed_value_is_caught(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    wb = load_workbook(out)
    wb[X.SHEET_DETAIL]["B2"] = "山田 太郎"
    wb.save(out)
    diffs = X.verify(out, TABLE, REASONS, JUDGMENTS)
    assert [(d.cell, d.what, d.original, d.readback) for d in diffs] == \
        [("B2", "値", "山田太郎", "山田 太郎")]


def test_an_emptied_reason_column_is_caught(tmp_path):
    """理由コードの列を人が消すと、塗りが残っていても突合が落ちる(色は理由の代わりにならない)。"""
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    wb = load_workbook(out)
    wb[X.SHEET_DETAIL]["H4"] = None
    wb.save(out)
    diffs = X.verify(out, TABLE, REASONS, JUDGMENTS)
    assert [(d.cell, d.what, d.original) for d in diffs] == [("H4", X.REASON_COLUMN, "値が無い")]
    assert len(X.read_fills(out)) == 14


def test_a_deleted_pending_sheet_is_caught(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    wb = load_workbook(out)
    del wb[X.SHEET_PENDING]
    wb.save(out)
    diffs = X.verify(out, TABLE, REASONS, JUDGMENTS)
    whats = {d.what for d in diffs}
    assert "座標 → 理由コード" in whats and "判断待ちの件数" in whats
    assert X.recover_reasons(out) == {}


def test_a_deleted_legend_sheet_is_caught(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    wb = load_workbook(out)
    del wb[X.SHEET_LEGEND]
    wb.save(out)
    assert [d.what for d in X.verify(out, TABLE, REASONS, JUDGMENTS)] == ["凡例の行数"]


def test_a_reason_code_outside_the_table_in_the_pending_sheet_is_caught(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    wb = load_workbook(out)
    wb[X.SHEET_PENDING]["F2"] = "なにか別の理由"
    wb.save(out)
    whats = [d.what for d in X.verify(out, TABLE, REASONS, JUDGMENTS)]
    assert "理由コードの表に無いコード" in whats


def test_write_deletes_the_working_file_when_the_readback_disagrees(tmp_path, monkeypatch):
    """書く段をわざと壊す。作業ファイルも出力ファイルも残らない。"""
    real = X._write_detail

    def broken(ws, table, reasons, judgments):
        real(ws, table, reasons, judgments)
        ws["B4"].fill = X._solid("EEF0F5")
    monkeypatch.setattr(X, "_write_detail", broken)
    r = X.write(tmp_path / "out.xlsx", TABLE, REASONS, JUDGMENTS)
    assert not r.written and not r.ok and r.diffs
    assert list(tmp_path.iterdir()) == []


def test_write_does_not_touch_an_existing_output_when_it_stops(tmp_path, monkeypatch):
    out = tmp_path / "out.xlsx"
    out.write_bytes(b"before")
    monkeypatch.setattr(X, "verify", lambda *a, **k: [X.Diff("明細", "B4", "塗り", "a", "b")])
    r = X.write(out, TABLE, REASONS, JUDGMENTS)
    assert not r.written and out.read_bytes() == b"before"
    assert [p.name for p in tmp_path.iterdir()] == ["out.xlsx"]


# ---- 色から理由を逆算する経路を作らない ------------------------------------------------------

def test_no_dict_in_the_module_is_keyed_by_a_colour():
    """塗り → 理由の対応表を持たないことを機械で確かめる(逆引きの経路を作らない)。"""
    hexish = re.compile(r"\A#?[0-9A-Fa-f]{6,8}\Z")
    for name, obj in vars(X).items():
        if isinstance(obj, dict):
            assert not [k for k in obj if isinstance(k, str) and hexish.match(k)], name


def test_the_module_has_no_function_that_takes_a_fill_and_returns_a_reason():
    names = [n for n in dir(X) if callable(getattr(X, n))]
    assert not [n for n in names if "reason_of_fill" in n or "fill_to" in n or "from_fill" in n]


def test_the_source_does_not_use_conditional_formatting_on_the_real_path():
    """条件付き書式は比較用の関数の中だけにあり、write の経路には無い。"""
    src = (ROOT / "excel_report.py").read_text(encoding="utf-8")
    body = src.split("def conditional_write")[0] + src.split("# ---- CLI")[1]
    assert "conditional_formatting" not in body


# ---- xlsx の中で値と色が別の場所に入っていること(zipfile で見る) ----------------------------

def test_the_reason_code_is_in_the_sheet_xml_and_not_in_the_styles(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    for code in REASONS.codes:
        where = X.where_it_lives(out, code)
        assert where == ["xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml",
                         "xl/worksheets/sheet3.xml"], code
        assert "xl/styles.xml" not in where


def test_the_fill_is_in_the_styles_and_not_in_the_detail_sheet_xml(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    for code in REASONS.codes:
        where = X.where_it_lives(out, REASONS.fill_of(code))
        assert where == ["xl/styles.xml", "xl/worksheets/sheet2.xml"], code
        assert "xl/worksheets/sheet1.xml" not in where     # 明細シートの XML に色は入らない


def test_the_xlsx_has_no_shared_strings_part(tmp_path):
    """この版の openpyxl は文字をシートの XML に直接書く(値の居場所がシート側だと確かめられる)。"""
    import zipfile
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    with zipfile.ZipFile(out) as z:
        assert "xl/sharedStrings.xml" not in z.namelist()


# ---- 比較用の 2 つの書き出し -----------------------------------------------------------------

def test_naive_write_paints_but_keeps_nothing_readable(tmp_path):
    out = tmp_path / "naive.xlsx"
    assert X.naive_write(out, TABLE, REASONS, JUDGMENTS) == 14
    assert len(X.read_fills(out)) == 14
    assert X.recover_reasons(out) == {}
    assert X.read_legend(out) == []
    assert load_workbook(out).sheetnames == [X.SHEET_DETAIL]
    assert X.read_values(out)[0] == list(TABLE.columns)       # 理由コードの列が無い


def test_conditional_write_leaves_the_colour_undetermined_in_the_file(tmp_path):
    out = tmp_path / "cond.xlsx"
    assert X.conditional_write(out, TABLE, REASONS, JUDGMENTS) == 5
    assert X.read_fills(out) == {}                            # ファイルの中に焼かれた塗りは無い
    assert X.read_values(out)[0][-1] == X.REASON_COLUMN        # 理由コードの列はある
    rules = load_workbook(out)[X.SHEET_DETAIL].conditional_formatting._cf_rules
    assert sum(len(v) for v in rules.values()) == 5


def test_verify_rejects_the_conditional_output(tmp_path):
    out = tmp_path / "cond.xlsx"
    X.conditional_write(out, TABLE, REASONS, JUDGMENTS)
    whats = {d.what for d in X.verify(out, TABLE, REASONS, JUDGMENTS)}
    assert "塗り" in whats and "座標 → 理由コード" in whats


# ---- 下流の工程(値だけを読む) ---------------------------------------------------------------

def copy_values_only(src: Path, dst: Path, sheet: str = X.SHEET_DETAIL) -> None:
    """値だけを別のブックに書き写す(値の貼り付け・再保存・別ツールでの書き出しに当たる)。"""
    wb = Workbook()
    wb.active.title = sheet
    for row in X.read_values(src, sheet):
        wb.active.append(row)
    wb.save(dst)


def to_csv_lines(src: Path, sheet: str = X.SHEET_DETAIL) -> list[str]:
    """CSV 化(値だけを読んでカンマでつなぐ)。"""
    return [",".join(r) for r in X.read_values(src, sheet)]


def test_values_only_copy_loses_every_fill(tmp_path):
    out, copied = tmp_path / "out.xlsx", tmp_path / "copied.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    copy_values_only(out, copied)
    assert len(X.read_fills(out)) == 14 and X.read_fills(copied) == {}


def test_values_only_copy_keeps_the_reason_column(tmp_path):
    out, copied = tmp_path / "out.xlsx", tmp_path / "copied.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    copy_values_only(out, copied)
    assert [r[-1] for r in X.read_values(copied)] == [r[-1] for r in X.read_values(out)]


def test_csv_of_the_pending_sheet_keeps_the_cell_coordinates(tmp_path):
    out = tmp_path / "out.xlsx"
    X.write(out, TABLE, REASONS, JUDGMENTS)
    lines = to_csv_lines(out, X.SHEET_PENDING)
    assert lines[1].startswith("明細,B4,3,氏名,")
    assert len(lines) - 1 == 14


def test_csv_of_a_colour_only_output_keeps_nothing(tmp_path):
    out = tmp_path / "naive.xlsx"
    X.naive_write(out, TABLE, REASONS, JUDGMENTS)
    lines = to_csv_lines(out)
    assert not [ln for ln in lines if any(c in ln for c in REASONS.codes)]


# ---- CLI --------------------------------------------------------------------------------------

def run_cli(args: list[str]) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, "-X", "utf8", str(ROOT / "excel_report.py")] + args,
                          capture_output=True, text=True, encoding="utf-8", errors="replace")


def test_cli_check_counts_by_reason():
    r = run_cli([str(FX / "sai_ichiran.csv"), "--judgments", str(FX / "hantei.csv"),
                 "--reasons", str(FX / "riyu_hyou.csv"), "--check"])
    assert r.returncode == 0
    assert "値が無い\t3" in r.stdout
    assert "判断待ち 14 件" in r.stderr


def test_cli_writes_and_reports(tmp_path):
    out = tmp_path / "out.xlsx"
    r = run_cli([str(FX / "sai_ichiran.csv"), str(out), "--judgments", str(FX / "hantei.csv"),
                 "--reasons", str(FX / "riyu_hyou.csv")])
    assert r.returncode == 0 and out.exists()
    assert "書いて読み戻して一致: 12 行 / 判断待ち 14 件" in r.stderr


def test_cli_naive_writes_colour_only(tmp_path):
    out = tmp_path / "naive.xlsx"
    r = run_cli([str(FX / "sai_ichiran.csv"), str(out), "--judgments", str(FX / "hantei.csv"),
                 "--reasons", str(FX / "riyu_hyou.csv"), "--naive"])
    assert r.returncode == 0 and X.recover_reasons(out) == {}


def test_cli_stops_on_a_reason_code_outside_the_table(tmp_path):
    bad = tmp_path / "hantei.csv"
    bad.write_text("行番号,列名,理由コード\n1,氏名,なにか別の理由\n", encoding="utf-8")
    r = run_cli([str(FX / "sai_ichiran.csv"), str(tmp_path / "o.xlsx"), "--judgments", str(bad),
                 "--reasons", str(FX / "riyu_hyou.csv")])
    assert r.returncode == 2 and "表に無いコード" in r.stderr


def test_cli_stops_on_a_broken_reason_table(tmp_path):
    bad = tmp_path / "riyu.csv"
    bad.write_text("理由コード,意味,塗り\n値が無い,意味,あか\n", encoding="utf-8")
    r = run_cli([str(FX / "sai_ichiran.csv"), str(tmp_path / "o.xlsx"),
                 "--judgments", str(FX / "hantei.csv"), "--reasons", str(bad)])
    assert r.returncode == 2 and "RRGGBB" in r.stderr


def test_cli_stops_when_the_row_number_is_not_a_number(tmp_path):
    bad = tmp_path / "hantei.csv"
    bad.write_text("行番号,列名,理由コード\n一,氏名,値が無い\n", encoding="utf-8")
    r = run_cli([str(FX / "sai_ichiran.csv"), str(tmp_path / "o.xlsx"), "--judgments", str(bad),
                 "--reasons", str(FX / "riyu_hyou.csv")])
    assert r.returncode == 2 and "行番号が数でない" in r.stderr


def test_cli_needs_a_destination():
    r = run_cli([str(FX / "sai_ichiran.csv"), "--judgments", str(FX / "hantei.csv"),
                 "--reasons", str(FX / "riyu_hyou.csv")])
    assert r.returncode == 2


# ---- 合成 fixture が宣言どおりであること -----------------------------------------------------

def test_fixture_rows_match_the_generator():
    assert len(MF.VALUE_ROWS) == 12 and len(MF.JUDGMENT_ROWS) == 14 and len(MF.REASON_ROWS) == 5
    assert TABLE.rows[0] == MF.VALUE_ROWS[0]


def test_fixture_has_no_real_data_marker():
    """氏名は架空の姓名の組み合わせだけ。fixture の生成器がそれを宣言している。"""
    doc = (FX / "make_fixtures.py").read_text(encoding="utf-8")
    assert "実データは 1 行も使っていない" in doc


# ---- 計測(記事の数値。固定 seed で決定論的に再現する) --------------------------------------

N = 1000
SEED = 20260927
PENDING_CELLS = 60
PENDING_ROWS = 55
DOUBLE_ROWS = 5
INJECT = (("値が無い", 14), ("表の外にある", 12), ("相手が決まらない", 11),
          ("形が両立しない", 16), ("往復で値が変わる", 7))
PRIMARY = (("値が無い", 12), ("表の外にある", 12), ("相手が決まらない", 11),
           ("形が両立しない", 13), ("往復で値が変わる", 7))
EXTRA = (("値が無い", 2), ("形が両立しない", 3))
COL_OF = {"値が無い": "氏名", "表の外にある": "項目", "相手が決まらない": "氏名",
          "形が両立しない": "前回", "往復で値が変わる": "氏名"}
EXTRA_COLUMN = "差"
SEI = ("山田", "鈴木", "佐藤", "髙橋", "渡辺", "伊藤", "中村", "小林", "加藤", "吉田")
MEI = ("太郎", "花子", "次郎", "一郎", "三郎", "四郎", "五郎", "六郎", "七郎", "八郎")
KOUMOKU = ("健康保険料", "厚生年金保険料", "介護保険料", "雇用保険料")
MEASURE_COLUMNS = ("社員番号", "氏名", "項目", "前回", "今回", "差")


def build_measure(seed: int = SEED) -> tuple[X.Table, list[X.Judgment]]:
    """差異一覧を 1,000 行合成し、上流の部品が返した判定に見立てて 60 セルに理由を付ける。"""
    rng = random.Random(seed)
    rows = []
    for i in range(N):
        zen = 15000 + rng.randrange(0, 400) * 50
        kon = zen + rng.choice([0, 0, 0, 525, -310, 1525])
        rows.append([f"{i + 1:06d}", rng.choice(SEI) + rng.choice(MEI), rng.choice(KOUMOKU),
                     str(zen), str(kon), str(kon - zen)])
    picks = rng.sample(range(1, N + 1), PENDING_ROWS)
    codes = [c for c, n in PRIMARY for _ in range(n)]
    judgments = [X.Judgment(picks[i], COL_OF[codes[i]], codes[i]) for i in range(PENDING_ROWS)]
    extra = [c for c, n in EXTRA for _ in range(n)]
    judgments += [X.Judgment(picks[i], EXTRA_COLUMN, extra[i]) for i in range(DOUBLE_ROWS)]
    index = {c: k for k, c in enumerate(MEASURE_COLUMNS)}
    for j in judgments:          # 理由に合わせて値の側も合成する(見た目が噛み合うように)
        cell = index[j.column]
        if j.code == "値が無い":
            rows[j.row - 1][cell] = ""
        elif j.code == "往復で値が変わる":
            rows[j.row - 1][cell] += "〜二郎"
        elif j.code == "形が両立しない":
            rows[j.row - 1][cell] = f"{int(rows[j.row - 1][cell] or 0):,}"
        elif j.code == "表の外にある":
            rows[j.row - 1][cell] += "(平成31年5月)"
    return X.Table(MEASURE_COLUMNS, tuple(tuple(r) for r in rows)), judgments


def test_measure_rows_are_built_as_declared():
    """合成の内訳を固定する(内訳が変われば記事の数字も変える)。"""
    table, judgments = build_measure()
    assert len(table.rows) == N
    assert len(judgments) == PENDING_CELLS == 60
    assert len({j.row for j in judgments}) == PENDING_ROWS == 55
    doubled = [r for r in {j.row for j in judgments} if len([j for j in judgments if j.row == r]) == 2]
    assert len(doubled) == DOUBLE_ROWS == 5
    assert X.counts_by_reason(REASONS, judgments) == dict(INJECT)
    X.check_judgments(table, REASONS, judgments)


def test_measure_this_tool_writes_and_the_readback_agrees(tmp_path):
    table, judgments = build_measure()
    out = tmp_path / "out.xlsx"
    r = X.write(out, table, REASONS, judgments)
    got = X.recover_reasons(out)
    print(f"\n[measure] this tool: {r.rows} rows, {r.pending} cells painted, "
          f"coordinates recovered {len(got)}, diffs {len(r.diffs)}, by reason {r.by_reason}")
    assert r.ok and r.rows == 1000 and r.pending == 60 and r.diffs == ()
    assert len(got) == 60 and got == X.expected_reason_map(table, judgments)
    assert len(X.read_fills(out)) == 60
    assert r.by_reason == dict(INJECT)


def test_measure_colour_only_output_and_the_downstream_steps(tmp_path):
    """色だけで示した出力は、値だけを読む工程を 1 つ通ると理由が 0 件になる。"""
    table, judgments = build_measure()
    naive, mine = tmp_path / "naive.xlsx", tmp_path / "mine.xlsx"
    X.naive_write(naive, table, REASONS, judgments)
    X.write(mine, table, REASONS, judgments)
    copied_naive, copied_mine = tmp_path / "naive2.xlsx", tmp_path / "mine2.xlsx"
    copy_values_only(naive, copied_naive)
    copy_values_only(mine, copied_mine)
    naive_rows = len([ln for ln in to_csv_lines(naive)[1:] if any(c in ln for c in REASONS.codes)])
    mine_rows = len([ln for ln in to_csv_lines(mine)[1:] if any(c in ln for c in REASONS.codes)])
    print(f"[measure] colour only: fills {len(X.read_fills(naive))}, "
          f"coordinates recovered {len(X.recover_reasons(naive))}, legend {len(X.read_legend(naive))}, "
          f"rows carrying a reason in the CSV {naive_rows}, fills after a values-only copy "
          f"{len(X.read_fills(copied_naive))}")
    print(f"[measure] this tool: fills {len(X.read_fills(mine))}, "
          f"coordinates recovered {len(X.recover_reasons(mine))}, legend {len(X.read_legend(mine))}, "
          f"rows carrying a reason in the CSV {mine_rows}, coordinates after a values-only copy "
          f"{len(X.recover_reasons(copied_mine))}, fills after a values-only copy "
          f"{len(X.read_fills(copied_mine))}")
    assert (len(X.read_fills(naive)), len(X.recover_reasons(naive)), len(X.read_legend(naive))) == (60, 0, 0)
    assert (len(X.read_fills(mine)), len(X.recover_reasons(mine)), len(X.read_legend(mine))) == (60, 60, 5)
    assert (naive_rows, mine_rows) == (0, 55)
    assert X.read_fills(copied_naive) == {} and X.read_fills(copied_mine) == {}
    assert len(X.recover_reasons(copied_mine)) == 0        # 明細だけ写すと座標の表は付いてこない


def test_measure_conditional_formatting_leaves_no_colour_in_the_file(tmp_path):
    table, judgments = build_measure()
    cond, mine = tmp_path / "cond.xlsx", tmp_path / "mine.xlsx"
    rules = X.conditional_write(cond, table, REASONS, judgments)
    X.write(mine, table, REASONS, judgments)
    print(f"[measure] conditional formatting: rules {rules}, fills readable from the file "
          f"{len(X.read_fills(cond))} / static {len(X.read_fills(mine))}")
    assert (rules, len(X.read_fills(cond)), len(X.read_fills(mine))) == (5, 0, 60)


def test_measure_where_the_values_and_the_fills_live(tmp_path):
    table, judgments = build_measure()
    out = tmp_path / "out.xlsx"
    X.write(out, table, REASONS, judgments)
    code_parts = {c: X.where_it_lives(out, c) for c in REASONS.codes}
    fill_parts = {c: X.where_it_lives(out, REASONS.fill_of(c)) for c in REASONS.codes}
    print(f"[measure] zipfile: a reason code is in {len(code_parts['値が無い'])} parts "
          f"{code_parts['値が無い']}, its fill is in {len(fill_parts['値が無い'])} parts "
          f"{fill_parts['値が無い']}")
    for c in REASONS.codes:
        assert len(code_parts[c]) == 3 and "xl/styles.xml" not in code_parts[c]
        assert fill_parts[c] == ["xl/styles.xml", "xl/worksheets/sheet2.xml"]


def test_measure_is_stable_across_runs():
    a_table, a_j = build_measure()
    b_table, b_j = build_measure()
    assert a_table == b_table and a_j == b_j
