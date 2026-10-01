"""excel_report の性質テスト(hypothesis)。README「いつも成り立つこと」の 1 文をそのまま検査にする。

いつも成り立つこと: 書き出しの結果は 2 通りしかない。
  (a) 出力ファイルができ、値だけを読み戻しても(色を見なくても)、全セルの値と、判定したセルごとの理由コードが
      入力のとおりに読み取れる。塗られているのは判定したセルだけで、塗りは理由コードの表のとおり
  (b) 出力ファイルはできず、止まった理由(受け取ったものの食い違い = SpecError、または読み戻しの差)が返る
表に無い理由コード・行・列や、同じセルへの 2 つ目の理由を渡すと、色を決める前に止まる(b の SpecError)。
どちらの場合も作業ファイルは残らない。

入力は乱数で作る。値の表(1〜4 列、0〜4 行)・理由コードの表(1〜3 個)・判定を、書ける値と噛み合った判定だけで
作り、半分の確率で値を 1 か所だけ困る値(= で始まる文字列、#N/A、前後の空白、改行・CR・タブ、32,768 字の
長い文字列、何でもありの文字列)に、半分の確率で判定を 1 件だけ困る形(表に無い行 0 / 行数 + 1、表に無い列、
表に無い理由コード、同じセルに 2 つ目)にする。値の表に `理由コード` 列を足すこともある。
derandomize=True で毎回同じ入力列を使う。database=None で見つけた例を保存しない(この版の hypothesis は実行した場所に .hypothesis/ を作るが、git と公開側への export では除外される)。

XML に書けない文字(タブ・改行・CR 以外の U+0000〜U+001F、U+FFFE、U+FFFF)を含む値も困る値に入れる。
README のとおり、書く前に食い違い(SpecError)として止まり、作業ファイルも残らないことを見る。

読み戻しは excel_report.verify / read_values / recover_reasons / read_fills を使わず、このテストの中で
openpyxl を直接開いて読む。期待するセル座標・理由コード・塗りもテストの中で別に作る。道具自身の突合が
壊れても、このテストで気づけるようにするため。
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from hypothesis import event, given, settings
from hypothesis import strategies as st
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT))
import excel_report as E  # noqa: E402

列の元 = ["社員番号", "氏名", "部署", "支給額"]
理由の元 = [("値が無い", "必要な値が空", "FFF1E8"), ("桁あふれ", "桁を超えた", "E8F0FF"),
           ("全角数字", "数字が全角", "F0FFE8")]                 # どのコードも他のコードの一部にならない
明細の見出しの右端 = "理由コード"
判断待ちの見出し = ["シート", "セル座標", "行番号", "列名", "値", "理由コード"]

書ける字 = list("あいうアイウ山田髙﨑㈱①ｱｲｳABCabc0123456789０１２-・ー ")
良い値 = st.one_of(st.just(""), st.from_regex(r"0[0-9]{1,5}", fullmatch=True),
                 st.text(alphabet=st.sampled_from(書ける字), min_size=1, max_size=10))
_書けない = {chr(c) for c in range(0x20) if chr(c) not in "\t\n\r"} | {"￾", "￿"}
_短い = st.text(alphabet=st.sampled_from(書ける字), max_size=4)
困る値 = st.one_of(
    st.sampled_from(["=1+1", "=A1", "#N/A", " 前後に空白 ", "改行\nあり", "CR\r\nLF", "\t",
                     "x" * 32768]),
    st.builds(lambda a, ch, b: a + ch + b, _短い, st.sampled_from(sorted(_書けない)), _短い),   # 書けない文字入り
    st.text(max_size=12),                                                                     # 何でもあり
)


@st.composite
def 入力(draw: st.DrawFn):
    columns = draw(st.lists(st.sampled_from(列の元), min_size=1, max_size=4, unique=True))
    n = draw(st.integers(0, 4))
    rows = [[draw(良い値) for _ in columns] for _ in range(n)]
    reasons = draw(st.lists(st.sampled_from(理由の元), min_size=1, max_size=3, unique=True))
    codes = [c for c, _, _ in reasons]
    judgments = []                                   # (行番号, 列名, 理由コード)。1 セル 1 件
    for i in range(1, n + 1):
        for c in columns:
            if draw(st.integers(0, 3)) == 0:
                judgments.append((i, c, draw(st.sampled_from(codes))))
    if rows and draw(st.booleans()):                 # 値を 1 か所だけ困る値に
        rows[draw(st.integers(0, n - 1))][draw(st.integers(0, len(columns) - 1))] = draw(困る値)
    if draw(st.booleans()):                          # 判定を 1 件だけ困る形に
        how = draw(st.sampled_from(["行 0", "行数 + 1", "表に無い列", "表に無い理由コード", "同じセルに 2 つ目",
                                    "理由コード列"]))
        if how == "行 0":
            judgments.append((0, columns[0], codes[0]))
        elif how == "行数 + 1":
            judgments.append((n + 1, columns[0], codes[0]))
        elif how == "表に無い列":
            judgments.append((max(n, 1), "表に無い列", codes[0]))
        elif how == "表に無い理由コード":
            judgments.append((max(n, 1), columns[0], "未登録"))
        elif how == "同じセルに 2 つ目" and judgments:
            r, c, _ = draw(st.sampled_from(judgments))
            judgments.append((r, c, draw(st.sampled_from(codes))))
        elif how == "理由コード列":
            columns = columns + [明細の見出しの右端]
            rows = [r + [""] for r in rows]
    return columns, rows, reasons, judgments


def 噛み合う(columns, rows, reasons, judgments) -> bool:
    """README「表に無い理由コード、表に無い行・列、同じセルに 2 つの理由、xlsx に書けない文字」と
    「理由コード列」が無いか。"""
    codes = {c for c, _, _ in reasons}
    cells = [(r, c) for r, c, _ in judgments]
    return (明細の見出しの右端 not in columns and len(set(cells)) == len(cells)
            and all(1 <= r <= len(rows) and c in columns and code in codes for r, c, code in judgments)
            and not any(set(v) & _書けない for row in rows for v in row))


def 座標(columns: list[str], r: int, c: str) -> str:
    """明細シートのセル座標。列は 4 つまでなので A〜D、見出しが 1 行あるので行は 1 つ下がる。"""
    return f"{'ABCD'[columns.index(c)]}{r + 1}"


def 文字(v) -> str:
    return "" if v is None else str(v)


@settings(max_examples=200, derandomize=True, database=None, deadline=None)
@given(入力=入力())
def test_いつも成り立つこと_値だけで理由が読めるファイルか_ファイル無しか(入力) -> None:
    columns, rows, reasons, judgments = 入力
    ok_in = 噛み合う(columns, rows, reasons, judgments)
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "out.xlsx"
        try:
            table = E.Table(tuple(columns), tuple(tuple(r) for r in rows))
            rt = E.ReasonTable(tuple(E.Reason(c, m, f) for c, m, f in reasons))
            r = E.write(out, table, rt, [E.Judgment(i, c, code) for i, c, code in judgments])
        except E.SpecError:
            event("(b) 食い違いで止まった")
            assert not ok_in, "噛み合っている入力が食い違いとして止まった"
            assert not out.exists()
        else:
            assert ok_in, "噛み合わない入力が、色を決める前に止まらなかった"
            if out.exists():
                event("(a) ファイルができた")
                assert r.ok and r.written and not r.diffs
                _確かめる(out, columns, rows, reasons, judgments)
            else:
                event("(b) 読み戻しの差で止まった")
                assert not r.written and r.diffs
        assert [p.name for p in Path(d).iterdir()] in ([], ["out.xlsx"]), "作業ファイルが残った"


def _確かめる(out: Path, columns, rows, reasons, judgments) -> None:
    """(a) の中身。値だけを読む(data_only=True)開き方と、塗りを読む開き方の 2 回で見る。"""
    fill = {c: f for c, _, f in reasons}
    wb = load_workbook(out, data_only=True)
    assert wb.sheetnames == ["明細", "凡例", "判断待ち"]

    明細 = [[文字(v) for v in row] for row in wb["明細"].iter_rows(values_only=True)]
    assert 明細[0] == columns + [明細の見出しの右端]
    assert len(明細) == 1 + len(rows)
    for i, (want, got) in enumerate(zip(rows, 明細[1:]), 1):
        assert got[:len(columns)] == want, f"{i} 行目の値が入力と違う"
        この行 = {code for r, _, code in judgments if r == i}
        理由の欄 = got[len(columns)]
        assert (理由の欄 == "") == (not この行)
        for code in fill:
            assert (code in 理由の欄) == (code in この行), f"{i} 行目の理由コード列に {code} が合わない"

    待ち = [[文字(v) for v in row] for row in wb["判断待ち"].iter_rows(values_only=True)]
    assert 待ち[0] == 判断待ちの見出し
    want = sorted(["明細", 座標(columns, r, c), str(r), c, rows[r - 1][columns.index(c)], code]
                  for r, c, code in judgments)
    assert sorted(待ち[1:]) == want, "判断待ちシートの「セル座標 → 理由コード」が入力と違う"

    塗り = {c.coordinate: str(c.fill.start_color.rgb)
           for row in load_workbook(out)["明細"].iter_rows() for c in row if c.fill.fill_type == "solid"}
    assert 塗り == {座標(columns, r, c): "FF" + fill[code] for r, c, code in judgments}, "塗りが判定と合わない"
